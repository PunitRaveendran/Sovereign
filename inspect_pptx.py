from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation('SIH2026-IDEA-Presentation-Format.pptx')
print(f'Slide count: {len(prs.slides)}')
print(f'Slide width: {prs.slide_width.inches:.2f} inches')
print(f'Slide height: {prs.slide_height.inches:.2f} inches')
print()

for i, slide in enumerate(prs.slides):
    print(f'=== SLIDE {i+1} ===')
    print(f'Layout: {slide.slide_layout.name}')
    for shape in slide.shapes:
        print(f'  Shape: [{shape.shape_type}] {shape.name!r} at ({shape.left/914400:.2f}", {shape.top/914400:.2f}") size=({shape.width/914400:.2f}" x {shape.height/914400:.2f}")')
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text:
                    print(f'    TEXT: {text[:120]}')
