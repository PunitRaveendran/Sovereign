import os
import sys
import glob
from pathlib import Path
import requests
from docx import Document
from PIL import Image, ImageDraw

def create_test_inspection_image(filename="grounded_test_inspection.png"):
    img = Image.new('RGB', (800, 320), color=(255, 255, 255))
    d = ImageDraw.Draw(img)
    content = """ASSET INSPECTION LOG: BALL VALVE BV-999
Date: 2026-09-04
Technician ID: TECH-409
Inspection Results:
- Wall Thickness: 91.2% (FAIL - minimum allowable: 95.0%)
- Actuator Torque: 68.5 N.m (FAIL - exceeds 55.0 N.m limit)
- Bonnet Flange Leakage: 7 drops/min observed
- Maximum Pitting Depth: 0.44mm (FAIL - permissible: 0.30mm)
Disposition: Quarantine asset immediately and flag for maintenance overhaul."""
    d.text((25, 25), content, fill=(0, 0, 0))
    img.save(filename)
    print(f"[TEST SETUP] Created synthetic inspection image: {filename}")
    return filename

def test_grounded_docx_generation():
    img_path = create_test_inspection_image()
    
    url = "http://127.0.0.1:8000/api/v1/task/upload"
    prompt = "Review the attached report and generate an approval note docx."
    
    print("\n--- STEP 1: Uploading image to Sovereign Multimodal Pipeline ---")
    with open(img_path, "rb") as f:
        files = {"file": (os.path.basename(img_path), f, "image/png")}
        data = {
            "user_id": "auditor_1",
            "role": "Engineer",
            "prompt": prompt
        }
        res = requests.post(url, files=files, data=data, timeout=60)
    
    print(f"Status Code: {res.status_code}")
    if res.status_code != 200:
        print(f"Error Response: {res.text}")
        sys.exit(1)
        
    res_data = res.json()
    print("API Response Summary:")
    print(f"- Status: {res_data.get('status')}")
    print(f"- Model Route: {res_data.get('routed_to_model')}")
    print(f"- Multimodal Meta: {res_data.get('multimodal_meta')}")
    
    agent_output = res_data.get("agent_output", {})
    observations = agent_output.get("observations", [])
    print(f"- Observations: {observations}")
    
    # Check generated docx
    output_dir = Path(r"c:\Users\Punit Raveendran\Desktop\SIH\manoj\outputs")
    docx_files = sorted(output_dir.glob("*.docx"), key=os.path.getmtime, reverse=True)
    if not docx_files:
        print("ERROR: No DOCX files found in output directory.")
        sys.exit(1)
        
    latest_docx = docx_files[0]
    print(f"\n--- STEP 2: Inspecting Generated DOCX ({latest_docx.name}) ---")
    doc = Document(str(latest_docx))
    
    all_text = []
    for p in doc.paragraphs:
        if p.text.strip():
            all_text.append(p.text.strip())
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                if cell.text.strip():
                    all_text.append(cell.text.strip())
                    
    full_doc_text = "\n".join(all_text)
    print("\n[DOCX Content Preview]:")
    print("-" * 50)
    print(full_doc_text[:1200])
    print("-" * 50)
    
    print("\n--- STEP 3: Grounding Verification Checks ---")
    expected_metrics = ["68.5 N.m", "7 drops/min", "0.44mm", "95.0%"]
    fake_defaults = ["96.5%", "48 N.m", "0.12mm"]
    
    grounded_passed = True
    for metric in expected_metrics:
        if metric in full_doc_text:
            print(f"  [PASS] Expected metric '{metric}' extracted from image and grounded in DOCX!")
        else:
            print(f"  [FAIL] Expected metric '{metric}' NOT found in DOCX!")
            grounded_passed = False
            
    for fake in fake_defaults:
        if fake in full_doc_text:
            print(f"  [FAIL] Hardcoded default metric '{fake}' WAS FOUND in DOCX! Deceptive fallback fired!")
            grounded_passed = False
        else:
            print(f"  [PASS] Hardcoded default metric '{fake}' is completely absent from DOCX.")
            
    if "PLACEHOLDER — NOT GROUNDED" in full_doc_text:
        print("  [FAIL] Unwanted 'PLACEHOLDER — NOT GROUNDED' found in DOCX despite valid image telemetry.")
        grounded_passed = False
    else:
        print("  [PASS] No ungrounded placeholders present; real telemetry was used.")
        
    print("\n--- STEP 4: Critique and Self-Correction Check ---")
    critique_output = agent_output.get("final_output", "")
    print(f"Critique Final Output:\n{critique_output.encode('ascii', errors='replace').decode('ascii')}")
    
    if grounded_passed:
        print("\n=======================================================")
        print("SUCCESS: End-to-end grounded document generation verified!")
        print("All extracted metrics trace directly to the image source.")
        print("=======================================================")
    else:
        print("\n=======================================================")
        print("FAILURE: Grounding verification failed.")
        print("=======================================================")
        sys.exit(1)

if __name__ == "__main__":
    test_grounded_docx_generation()
